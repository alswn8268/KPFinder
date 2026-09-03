from app.report import build_excel_report, build_final_state, render_current_tree, render_proposed_tree
from app.scanner import FileEntry
from datetime import datetime


def _entry(rel_path: str, status: str = "ok") -> FileEntry:
    e = FileEntry(
        path=rel_path,
        relative_path=rel_path,
        name=rel_path.split("/")[-1],
        ext=".txt",
        size=10,
        modified=datetime.now(),
    )
    e.summary_status = status
    return e


def test_render_current_tree_groups_by_folder():
    entries = [_entry("a.txt"), _entry("사업계획/b.txt"), _entry("사업계획/c.txt")]
    tree_text = render_current_tree(entries)

    assert "📁 사업계획/" in tree_text
    assert "📄 a.txt" in tree_text
    assert "📄 b.txt" in tree_text


def test_render_proposed_tree_empty_assignments():
    assert render_proposed_tree({}) == "(제안된 이동이 없습니다)"


def test_render_proposed_tree_with_assignments():
    text = render_proposed_tree({"a.txt": "보고서/a.txt", "b.txt": "보고서/b.txt"})
    assert "📁 보고서/" in text
    assert "📄 a.txt" in text


def test_build_final_state_covers_moved_kept_excluded_unsupported():
    entries = [
        _entry("moved.txt"),
        _entry("kept.txt"),
        _entry("excluded.txt"),
        _entry("bad.zip", status="unsupported"),
    ]
    valid_assignments = {"moved.txt": {"dst": "보고서/moved.txt", "reason": "", "confidence": "높음"}}
    rows = build_final_state(entries, valid_assignments, excluded_srcs={"excluded.txt"}, rejected=[])
    by_path = {r["경로"]: r["상태"] for r in rows}

    assert by_path["moved.txt"] == "이동"
    assert by_path["kept.txt"] == "유지"
    assert by_path["excluded.txt"] == "제외"
    assert by_path["bad.zip"] == "검토 필요"


def test_build_final_state_flags_destination_conflict():
    entries = [_entry("a.txt"), _entry("b.txt")]
    valid_assignments = {
        "a.txt": {"dst": "보고서/same.txt", "reason": "", "confidence": "높음"},
        "b.txt": {"dst": "보고서/same.txt", "reason": "", "confidence": "높음"},
    }
    rows = build_final_state(entries, valid_assignments, excluded_srcs=set(), rejected=[])
    assert all(r["상태"] == "충돌" for r in rows)


def test_build_excel_report_produces_readable_workbook(tmp_path):
    from openpyxl import load_workbook

    entries = [_entry("a.txt")]
    final_state = build_final_state(entries, {}, excluded_srcs=set(), rejected=[])
    data = build_excel_report(entries, {}, final_state)

    out_path = tmp_path / "report.xlsx"
    out_path.write_bytes(data)
    wb = load_workbook(str(out_path))
    assert "파일별 최종 상태" in wb.sheetnames
