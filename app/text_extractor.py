"""지원 형식에서 텍스트를 추출한다.

지원: .txt/.docx/.xlsx/.csv/.pptx/.pdf/.hwpx/.hwp

.hwp(바이너리 OLE 복합 문서)는 olefile + 레코드 파싱으로 본문 텍스트를 최대한
읽어내되(부분 지원), 문서 구조가 예상과 다르거나 파싱이 실패하면 HwpParseError를
던져 호출자가 파일명 기반 분류로 넘어가도록 한다("완전 지원이 어렵더라도 최소
대응은 한다"는 기획서 §13.2 원칙).
"""

import csv
import re
import struct
import zipfile
from xml.etree import ElementTree

from app.scanner import SUPPORTED_TEXT_EXTS


class ExtractionError(Exception):
    """텍스트 추출 중 발생한 오류(일반)."""


class EncryptedDocumentError(ExtractionError):
    """문서가 암호화되어 있어 본문을 읽을 수 없는 경우."""


class OcrNeededError(ExtractionError):
    """이미지 형태의 PDF 등, OCR 없이는 본문을 읽을 수 없는 경우."""


class HwpParseError(ExtractionError):
    """HWP 바이너리 구조 파싱에 실패해 본문을 읽지 못한 경우(파일명 기반 분류로 대체)."""


def is_supported(ext: str) -> bool:
    return ext.lower() in SUPPORTED_TEXT_EXTS


def extract_text(path: str, ext: str, max_chars: int = 4000) -> str:
    ext = ext.lower()
    try:
        if ext == ".txt":
            return _extract_txt(path, max_chars)
        if ext == ".docx":
            return _extract_docx(path, max_chars)
        if ext == ".csv":
            return _extract_csv(path, max_chars)
        if ext == ".xlsx":
            return _extract_xlsx(path, max_chars)
        if ext == ".pptx":
            return _extract_pptx(path, max_chars)
        if ext == ".pdf":
            return _extract_pdf(path, max_chars)
        if ext == ".hwpx":
            return _extract_hwpx(path, max_chars)
        if ext == ".hwp":
            return _extract_hwp(path, max_chars)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - 어떤 라이브러리든 실패는 추출 오류로 통일
        raise ExtractionError(f"{ext} 텍스트 추출 실패: {exc}") from exc
    raise ExtractionError(f"지원하지 않는 확장자: {ext}")


def _extract_txt(path: str, max_chars: int) -> str:
    for encoding in ("utf-8", "cp949", "euc-kr"):
        try:
            with open(path, "r", encoding=encoding) as f:
                return f.read(max_chars)
        except (UnicodeDecodeError, LookupError):
            continue
    with open(path, "rb") as f:
        return f.read(max_chars).decode("utf-8", errors="ignore")


def _extract_docx(path: str, max_chars: int) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(parts)[:max_chars]


def _extract_csv(path: str, max_chars: int) -> str:
    for encoding in ("utf-8", "cp949", "euc-kr"):
        try:
            lines = []
            total = 0
            with open(path, "r", encoding=encoding, newline="") as f:
                for row in csv.reader(f):
                    line = ", ".join(row)
                    lines.append(line)
                    total += len(line)
                    if total >= max_chars:
                        break
            return "\n".join(lines)[:max_chars]
        except (UnicodeDecodeError, LookupError):
            continue
    return ""


def _extract_xlsx(path: str, max_chars: int) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    chunks = []
    total = 0
    for sheet in wb.worksheets:
        chunks.append(f"[시트: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if not values:
                continue
            line = ", ".join(values)
            chunks.append(line)
            total += len(line)
            if total >= max_chars:
                break
        if total >= max_chars:
            break
    return "\n".join(chunks)[:max_chars]


def _extract_pptx(path: str, max_chars: int) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    chunks = []
    total = 0
    for i, slide in enumerate(prs.slides, start=1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if texts:
            chunk = f"[슬라이드 {i}] " + " / ".join(texts)
            chunks.append(chunk)
            total += len(chunk)
        if total >= max_chars:
            break
    return "\n".join(chunks)[:max_chars]


def _extract_pdf(path: str, max_chars: int) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            # 빈 암호로 열리는 PDF도 있으므로 한 번 시도는 해 본다.
            if reader.decrypt("") == 0:
                raise EncryptedDocumentError("암호로 보호된 PDF입니다.")
        except Exception as exc:
            if isinstance(exc, EncryptedDocumentError):
                raise
            raise EncryptedDocumentError("암호로 보호된 PDF입니다.") from exc

    chunks = []
    total = 0
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            chunks.append(text)
            total += len(text)
        if total >= max_chars:
            break

    if not chunks:
        raise OcrNeededError(
            "이 PDF는 이미지 형태로 저장되어 본문을 읽지 못했습니다. "
            "파일명과 기존 폴더를 기준으로만 임시 분류합니다."
        )
    return "\n".join(chunks)[:max_chars]


_HWPX_SECTION_RE = re.compile(r"^Contents/section\d+\.xml$")


def _extract_hwpx(path: str, max_chars: int) -> str:
    """HWPX(OWPML, ZIP+XML 기반)에서 본문 텍스트를 추출한다."""
    chunks = []
    total = 0
    with zipfile.ZipFile(path) as zf:
        section_names = sorted(n for n in zf.namelist() if _HWPX_SECTION_RE.match(n))
        if not section_names:
            raise ExtractionError("HWPX 본문 섹션을 찾을 수 없습니다.")
        for name in section_names:
            with zf.open(name) as f:
                tree = ElementTree.parse(f)
            texts = [t.strip() for t in tree.getroot().itertext() if t and t.strip()]
            chunk = " ".join(texts)
            if chunk:
                chunks.append(chunk)
                total += len(chunk)
            if total >= max_chars:
                break
    return "\n".join(chunks)[:max_chars]


HWPTAG_PARA_TEXT = 0x43  # HWP5 바이너리 레코드 태그: 문단 텍스트
_CONTROL_CHAR_RE = re.compile(r"[\x00-\x08\x0b-\x1f]")


def _extract_hwp(path: str, max_chars: int) -> str:
    """HWP 5.0(OLE 복합 문서) 본문을 최대한 읽어낸다(부분 지원, best-effort).

    FileHeader의 압축/암호화 플래그를 확인하고, BodyText/Section* 스트림을
    필요 시 raw deflate로 해제한 뒤 레코드를 순회하며 HWPTAG_PARA_TEXT(0x43)
    레코드만 UTF-16LE로 디코드한다. 인라인 컨트롤 문자(필드/표/그림 등)까지
    완벽히 걸러내지는 못해 약간의 잡음이 섞일 수 있지만, 요약·분류 힌트로는
    충분한 수준의 본문을 얻을 수 있다. 구조를 해석할 수 없으면 HwpParseError를
    던져 호출자가 파일명 기반 분류로 대체하도록 한다.
    """
    import zlib

    try:
        import olefile
    except ImportError as exc:
        raise HwpParseError("HWP 파싱에 필요한 olefile 라이브러리가 설치되어 있지 않습니다.") from exc

    try:
        ole = olefile.OleFileIO(path)
    except Exception as exc:
        raise HwpParseError(f"HWP 파일 구조(OLE)를 열 수 없습니다: {exc}") from exc

    try:
        if not ole.exists("FileHeader"):
            raise HwpParseError("HWP FileHeader를 찾을 수 없습니다.")
        header = ole.openstream("FileHeader").read()
        if len(header) < 40 or not header.startswith(b"HWP Document File"):
            raise HwpParseError("HWP 시그니처가 올바르지 않습니다.")
        flags = struct.unpack("<I", header[36:40])[0]
        compressed = bool(flags & 0x01)
        encrypted = bool(flags & 0x02)
        if encrypted:
            raise EncryptedDocumentError("암호로 보호된 HWP 문서입니다.")

        section_entries = sorted(
            (entry for entry in ole.listdir() if len(entry) == 2 and entry[0] == "BodyText"),
            key=lambda e: e[1],
        )
        if not section_entries:
            raise HwpParseError("HWP BodyText 섹션을 찾을 수 없습니다.")

        chunks = []
        total = 0
        for entry in section_entries:
            raw = ole.openstream(entry).read()
            data = zlib.decompressobj(-15).decompress(raw) if compressed else raw
            text = "".join(_iter_para_texts(data))
            text = _CONTROL_CHAR_RE.sub(" ", text).strip()
            if text:
                chunks.append(text)
                total += len(text)
            if total >= max_chars:
                break

        if not chunks:
            raise HwpParseError("HWP 본문에서 텍스트를 찾지 못했습니다.")
        return "\n".join(chunks)[:max_chars]
    finally:
        ole.close()


def _iter_para_texts(data: bytes):
    """압축 해제된 HWP 레코드 스트림에서 문단 텍스트 레코드만 뽑아 디코드한다."""
    pos = 0
    n = len(data)
    while pos + 4 <= n:
        header = struct.unpack_from("<I", data, pos)[0]
        tag_id = header & 0x3FF
        size = (header >> 20) & 0xFFF
        pos += 4
        if size == 0xFFF:
            if pos + 4 > n:
                break
            size = struct.unpack_from("<I", data, pos)[0]
            pos += 4
        body = data[pos : pos + size]
        pos += size
        if tag_id == HWPTAG_PARA_TEXT and body:
            try:
                yield body.decode("utf-16le", errors="ignore")
            except Exception:  # noqa: BLE001 - 개별 레코드 디코드 실패는 건너뛴다
                continue
